"""Exports of an immutable calculation; user text is escaped in both formats."""
import io
import json
from xml.sax.saxutils import escape
from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from .finance import FIELDS
from .domain import OPERATIONS
from .templatetags.workflow_tags import status


def sections(c):
    result=c.result
    rows=[('Расчёт', [['Показатель','Значение'],['Проект',c.version.study.name],['Версия требований',c.version.number],['Конфигурация',c.pk],['Дата',c.created_at.isoformat()],['Статус',result['status']],['Горизонт','5 лет, постоянные цены; аренда без выкупа']])]
    rows.append(('Экономика', [['Вариант','CAPEX, руб.','OPEX в год','TCO, 5 лет','NPV','Окупаемость, лет']]+[[s['name'],float(s['capex']),float(s['opex']),float(s['tco']),float(s['npv']),s['payback']] for s in result['scenarios']]))
    rows.append(('Параметры', [['Параметр','Значение']]+[[FIELDS[k][0],v] for k,v in c.version.finance.items()]))
    rows.append(('Операции', [['ID','Название','После','Тип','Дистанция, м','Буфер','Мест']]+[[o['id'],o['name'],', '.join(o['predecessors']),o['code'],o['distance_m'],o['buffer'],o['stations']] for o in c.version.requirements['operations']]))
    rows.append(('Состав', [['Группа','Исполнитель','Активных','С резервом']]+[[g['id'],g.get('data',{}).get('model','Человек'),g['count'],g.get('total_count',g['count'])] for g in result['groups']]))
    rows.append(('Денежные потоки', [['Вариант','Год','Эффект','Накопленный эффект','Дисконтированный эффект']]+[[s['name'],f['year'],float(f['flow']),float(f['cumulative']),float(f['discounted'])] for s in result['scenarios'] for f in s['flows']]))
    sim=result.get('simulation',{})
    notes=result['errors']+result['missing']+result.get('preliminary',[])+sim.get('assumptions',[])
    rows.append(('Проверка', [['Показатель','Значение'],['Результат',sim.get('reason','Расчёт остановлен')],['Пиковый поток',c.version.requirements['peak_hour']],['Поток в прогоне',sim.get('throughput')],['Проверенная нагрузка',result.get('load_limit',{}).get('tested_stable')],*[["Примечание",n] for n in notes]]))
    rows.append(('Источники', [['Версия','Модель','Источник','Комплектация']]+[[r['revision'],r['data']['model'],r['data'].get('source',''),r['data'].get('configuration','')] for r in c.snapshot.values()]))
    rows[0][1][-1][1]=f"{c.version.finance.get('horizon',5)} лет, постоянные цены; аренда без выкупа"
    rows[0][1][4][1]=timezone.localtime(c.created_at).strftime('%d.%m.%Y %H:%M %Z')
    rows[0][1].append(['Версия модели',result.get('model_version','workflow-2.1')])
    rows[0][1].append(['Заключение',result.get('conclusion','Экономика недоступна до проверки выполнимости и исходных данных')])
    rows[1][1][0][3]='TCO за горизонт'
    rows.append(('Показатели', [['Вариант','Эффект в год','ROI по ТЗ, %','Чистый ROI, %','Цена операции']]+[[s['name'],float(s['annual_effect']),s['roi'],s.get('net_roi'),float(s['cost_per_operation'])] for s in result['scenarios']]))
    rows.append(('Статьи затрат', [['Вариант','Тип','Статья','Сумма, руб.']]+[[s['name'],kind,key,float(value)] for s in result['scenarios'] for kind in ['capex_items','opex_items'] for key,value in s.get(kind,{}).items()]))
    rows.append(('Чувствительность', [['Параметр','Изменение, %','Статус','Сценарий','TCO','Эффект в год']]+[[r['parameter'],r['delta'],r['status'],s['name'],float(s['tco']),float(s['annual_effect'])] for r in result.get('sensitivity',[]) for s in r['scenarios']]+[[r['parameter'],r['delta'],r['status'],r['reason'],None,None] for r in result.get('sensitivity',[]) if not r['scenarios']]))
    labels={'area':'Площадь склада, м²','active_area':'Рабочая зона, м²','aisle_m':'Минимальный проход, м','payload_kg':'Масса паллеты, кг','temperature':'Температура, °C','peak_hour':'Пиковый поток, паллет/ч','daily_moves':'Паллет в сутки','clearance_m':'Боковой зазор, м','allow_human':'Ручные этапы разрешены'}
    rows.append(('Объект', [['Параметр','Значение']]+[[labels.get(k,k),v] for k,v in c.version.requirements.items() if not isinstance(v,(list,dict))]))
    rows.append(('Формулы', [['Показатель','Метод'],['TCO','CAPEX + OPEX × горизонт'],['Годовой эффект','OPEX текущего процесса минус OPEX варианта'],['ROI по ТЗ','Накопленный эффект / CAPEX × 100%'],['Чистый ROI','(Накопленный эффект минус CAPEX) / CAPEX × 100%'],['NPV','Сумма дисконтированных эффектов минус CAPEX'],['Окупаемость','CAPEX / положительный годовой эффект'],['Допущения','Собственные средства; постоянные цены с НДС; без налоговых эффектов и ликвидационной стоимости. CAPEX учитывается один раз; амортизация повторно не вычитается. Аренда без выкупа. Резерв парка 10%.']]))
    for _,table in rows:
        for row in table:
            for i,value in enumerate(row):
                if isinstance(value,bool):row[i]='Да' if value else 'Нет'
                elif isinstance(value,str):row[i]={'capex_items':'CAPEX','opex_items':'OPEX за год'}.get(value,OPERATIONS.get(value,status(value)))
    return rows


def xlsx(c):
    book=Workbook();book.remove(book.active)
    for title,rows in sections(c):
        sheet=book.create_sheet(title)
        for row in rows:
            sheet.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='176C64')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=38
            for cell in col:cell.alignment=Alignment(wrap_text=True,vertical='top')
    out=io.BytesIO();book.save(out);return out.getvalue()


def pdf(c):
    if 'WorkflowRu' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('WorkflowRu',str(settings.BASE_DIR/'static/fonts/DejaVuSans.ttf')))
    body=ParagraphStyle('body',fontName='WorkflowRu',fontSize=8,leading=11,spaceAfter=6)
    title=ParagraphStyle('title',parent=body,fontSize=17,leading=22,spaceAfter=14)
    sub=ParagraphStyle('sub',parent=body,fontSize=12,leading=16,spaceBefore=12,spaceAfter=8,keepWithNext=True)
    story=[Paragraph('Робоскоп. Оценка автоматизации склада',title),Paragraph('Предварительная оценка. Условия требуют согласования с поставщиками. Реальные договоры и платежи в прототипе не выполняются.',body)]
    for heading,rows in sections(c):
        story.append(Paragraph(heading,sub))
        formatted=[[Paragraph(escape('Не применяется' if v is None else f'{v:,.2f}'.replace(',',' ').replace('.',',') if isinstance(v,float) else str(v)),body) for v in row] for row in rows]
        table=Table(formatted,colWidths=[510/len(rows[0])]*len(rows[0]),repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e4f1ef')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.25,colors.HexColor('#d9e1e0')),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4)]))
        story.extend([table,Spacer(1,8)])
    def footer(canvas,doc):
        canvas.setFont('WorkflowRu',8);canvas.drawString(42,20,'Робоскоп · предварительная оценка');canvas.drawRightString(552,20,str(doc.page))
    out=io.BytesIO();SimpleDocTemplate(out,leftMargin=42,rightMargin=42,topMargin=36,bottomMargin=36,title='Робоскоп: расчёт конфигурации').build(story,onFirstPage=footer,onLaterPages=footer);return out.getvalue()
