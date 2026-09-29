import csv
import hashlib
import io
import json
import zipfile
from datetime import timedelta
from functools import wraps
from django.conf import settings
from django.contrib.auth import authenticate, login, logout, get_user_model
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.db import transaction, connection, IntegrityError
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, get_object_or_404
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_http_methods
from openpyxl import load_workbook, Workbook
from .engine import validate, defaults, calculate, applicability, product_dict, InputError
from .models import Product, Normative, Project, ProjectVersion, SimulationRun, LoginAttempt, ServiceHeartbeat
from .schema import FIELDS, MODEL_VERSION, DATA_VERSION, FORMULAS
from .catalog import import_catalog
from .reports import excel_report, pdf_report
from uuid import UUID

def api(methods=('GET',),auth=False,staff=False):
    def decorate(fn):
        @wraps(fn)
        @require_http_methods(methods)
        def wrapped(request,*args,**kwargs):
            if auth and not request.user.is_authenticated:return JsonResponse({'error':'Войдите, чтобы сохранить или открыть свой проект'},status=401)
            if staff and not request.user.is_staff:return JsonResponse({'error':'Действие доступно администратору'},status=403)
            try:return fn(request,*args,**kwargs)
            except InputError as e:return JsonResponse({'error':str(e),'fields':e.errors},status=400)
            except (ValueError,TypeError,KeyError,ValidationError) as e:
                return JsonResponse({'error': '; '.join(e.messages) if isinstance(e,ValidationError) else str(e)},status=400)
        return wrapped
    return decorate

def body(request):
    try:data=json.loads(request.body or b'{}')
    except (ValueError,UnicodeDecodeError):raise ValueError('Не удалось прочитать JSON')
    if not isinstance(data,dict):raise ValueError('Нужен JSON-объект')
    return data

def get_calculation(data,user):
    if data.get('project_id'):
        if not user.is_authenticated:raise ValueError('Войдите для доступа к проекту')
        project=get_object_or_404(Project,id=data['project_id'],owner=user)
        version=get_object_or_404(project.versions,revision=int(data.get('calculation_revision',data.get('revision',project.revision))))
        return version.result
    if data.get('object_type','warehouse')!='warehouse':raise ValueError('Полный расчёт реализован для склада. Для других объектов доступен каталог и исходные параметры.')
    v=validate(data.get('inputs',{}),defaults())
    if not data.get('product_id'): raise ValueError('Выберите продукт для расчёта')
    try: product_id=UUID(str(data['product_id']))
    except (ValueError,TypeError):raise ValueError('Некорректный идентификатор продукта')
    p=get_object_or_404(Product,id=product_id)
    result=calculate(v,product_dict(p))
    norms={n.key:n for n in Normative.objects.all()}
    result['assumptions']={f['key']:dict(source=norms[f['key']].source if f['key'] in norms else f['source'],
        default=norms[f['key']].value if f['key'] in norms else f['default'],overridden=v[f['key']]!=(norms[f['key']].value if f['key'] in norms else f['default'])) for f in FIELDS}
    return result

@ensure_csrf_cookie
def index(request):return render(request,'index.html')

@api()
def bootstrap(request):
    datasets=json.loads((settings.BASE_DIR/'data/objects.json').read_text())
    fields=[dict(f) for f in FIELDS];norms={n.key:n for n in Normative.objects.all()}
    for f in fields:
        if f['key'] in norms:f.update(default=norms[f['key']].value,source=norms[f['key']].source)
    return JsonResponse(dict(user=dict(username=request.user.username,admin=request.user.is_staff) if request.user.is_authenticated else None,
        passwordless_demo=getattr(request,"passwordless_demo",False),fields=fields,defaults=defaults(),objects=datasets,model_version=MODEL_VERSION,data_version=DATA_VERSION,
        formulas=FORMULAS,product_count=Product.objects.count()))

@api()
def catalog(request):
    values=validate(json.loads(request.GET.get('inputs','{}')),defaults())
    object_type=request.GET.get('object','warehouse')
    if object_type not in ['warehouse','airport','hospital']:raise ValueError('Неизвестный тип объекта')
    data=[]
    for p in Product.objects.all():
        item=product_dict(p);item['match']=applicability(item,values,object_type);data.append(item)
    data.sort(key=lambda p:(p['match']['state']=='excluded',-p['match']['score'],p['name']))
    return JsonResponse(dict(products=data))

@api(('POST',))
def calculate_view(request):return JsonResponse(get_calculation(body(request),request.user))

@api(('POST',))
def auth_view(request,action):
    if action not in ('login','logout'):return JsonResponse({'error':'Регистрация временно недоступна'},status=404)
    if action=='logout':logout(request);return JsonResponse({'ok':True})
    previous_session=request.session.session_key
    data=body(request);username=str(data.get('username','')).strip();password=data.get('password','')
    if not isinstance(password,str):raise ValueError('Некорректный пароль')
    if not username or len(username)>80:raise ValueError('Имя пользователя: 1–80 символов')
    key=hashlib.sha256((request.META.get('REMOTE_ADDR','')+'|'+username.lower()).encode()).hexdigest()
    attempt,_=LoginAttempt.objects.get_or_create(key=key)
    if attempt.failures>=8 and timezone.now()-attempt.updated_at<timedelta(minutes=15):
        return JsonResponse({'error':'Слишком много попыток. Повторите через 15 минут.'},status=429)
    if action=='login':
        u=authenticate(request,username=username,password=password)
        if not u:
            if timezone.now()-attempt.updated_at>=timedelta(minutes=15):attempt.failures=0
            attempt.failures+=1;attempt.save()
            return JsonResponse({'error':'Неверное имя пользователя или пароль'},status=400)
        if data.get('role')=='admin' and not u.is_staff:
            return JsonResponse({'error':'У аккаунта нет прав администратора'},status=403)
        login(request,u)
    else:raise ValueError('Неизвестное действие')
    if previous_session:
        SimulationRun.objects.filter(owner=None,session_key=previous_session).update(owner=u)
    attempt.delete();return JsonResponse({'username':u.username,'admin':u.is_staff})

def project_summary(p):
    return dict(id=str(p.id),name=p.name,object_type=p.object_type,revision=p.revision,updated_at=p.updated_at.isoformat(),
                robots=p.result.get('fleet',{}).get('robots'))

def attach_simulation(result,data,request):
    if not data.get('simulation_id'):return result
    qs=SimulationRun.objects.filter(owner=request.user) if request.user.is_authenticated else SimulationRun.objects.filter(owner=None,session_key=request.session.session_key or 'invalid')
    run=get_object_or_404(qs,id=data['simulation_id'],status='done')
    original=run.payload['calculation']
    if original['inputs']!=result['inputs'] or original['product']!=result['product'] or original['model_version']!=result['model_version']:
        raise ValueError('Параметры или каталог изменились после симуляции. Выполните новый прогон.')
    result['simulation_summary']={k:v for k,v in run.result.items() if k not in ['events','arrivals','completions','per_robot']}
    return result

@api(('GET','POST'),auth=True)
def projects(request):
    if request.method=='GET':return JsonResponse({'projects':[project_summary(p) for p in Project.objects.filter(owner=request.user)]})
    data=body(request);name=str(data.get('name','Новый проект')).strip()
    if not name or len(name)>160:raise ValueError('Название проекта должно содержать 1–160 символов')
    result=attach_simulation(get_calculation(data,request.user),data,request)
    with transaction.atomic():
        p=Project.objects.create(owner=request.user,name=name,payload=dict(inputs=result['inputs'],product_id=result['product']['id'],object_type='warehouse'),result=result)
        ProjectVersion.objects.create(project=p,revision=1,payload=p.payload,result=result)
    return JsonResponse(project_summary(p),status=201)

@api(('GET','PUT','DELETE'),auth=True)
def project_detail(request,id):
    p=get_object_or_404(Project,id=id,owner=request.user)
    if request.method=='DELETE':p.delete();return JsonResponse({'ok':True})
    if request.method=='GET':
        revision=int(request.GET.get('revision',p.revision));version=get_object_or_404(p.versions,revision=revision)
        return JsonResponse(dict(**project_summary(p),loaded_revision=version.revision,payload=version.payload,result=version.result,
             versions=[dict(revision=v.revision,date=v.created_at.isoformat()) for v in p.versions.all()]))
    data=body(request)
    result=attach_simulation(get_calculation(data,request.user),data,request)
    with transaction.atomic():
        p=Project.objects.select_for_update().get(id=p.id)
        if int(data.get('revision',0))!=p.revision:return JsonResponse({'error':'Проект изменён в другой вкладке. Откройте актуальную версию.'},status=409)
        name=str(data.get('name',p.name)).strip()
        if not name or len(name)>160:raise ValueError('Название: 1–160 символов')
        p.name=name;p.revision+=1;p.payload=dict(inputs=result['inputs'],product_id=result['product']['id'],object_type='warehouse');p.result=result;p.save()
        ProjectVersion.objects.create(project=p,revision=p.revision,payload=p.payload,result=result)
    return JsonResponse(project_summary(p))

@api(('POST',),auth=True)
def project_copy(request,id):
    p=get_object_or_404(Project,id=id,owner=request.user)
    with transaction.atomic():
        new=Project.objects.create(owner=request.user,name=(p.name+' · копия')[:160],payload=p.payload,result=p.result)
        ProjectVersion.objects.create(project=new,revision=1,payload=new.payload,result=new.result)
    return JsonResponse(project_summary(new),status=201)

@api(('POST',))
def simulation_start(request):
    data=body(request);result=get_calculation(data,request.user)
    hours=float(data.get('hours',4));seed=int(data.get('seed',42))
    if not .25<=hours<=8:raise ValueError('Горизонт симуляции: 0,25–8 часов')
    if not 0<=seed<=2147483647:raise ValueError('Seed: от 0 до 2147483647')
    if result['fleet']['robots']>80 or result['fleet']['peak_demand']*hours>20000:raise ValueError('Сократите область симуляции: максимум 80 роботов и 20 000 операций')
    if not request.session.session_key:request.session.create()
    owner=request.user if request.user.is_authenticated else None
    mine=SimulationRun.objects.filter(owner=owner) if owner else SimulationRun.objects.filter(session_key=request.session.session_key)
    if mine.filter(status__in=['pending','running'],created_at__gt=timezone.now()-timedelta(minutes=3)).exists():
        return JsonResponse({'error':'Предыдущий прогон ещё выполняется'},status=409)
    if SimulationRun.objects.filter(status__in=['pending','running']).count()>=50:
        return JsonResponse({'error':'Очередь симуляций заполнена. Повторите позже.'},status=429)
    run=SimulationRun.objects.create(owner=owner,session_key=request.session.session_key,
          payload=dict(calculation=result,hours=hours,seed=seed))
    return JsonResponse({'id':str(run.id),'status':run.status},status=202)

@api()
def simulation_status(request,id):
    qs=SimulationRun.objects.filter(owner=request.user) if request.user.is_authenticated else SimulationRun.objects.filter(owner=None,session_key=request.session.session_key or 'invalid')
    run=get_object_or_404(qs,id=id)
    return JsonResponse(dict(id=str(run.id),status=run.status,error=run.error,result=run.result if run.status=='done' else None))

@api(('POST',))
def export_report(request,format):
    data=body(request)
    result=attach_simulation(get_calculation(data,request.user),data,request)
    if format=='pdf':content=pdf_report(result);mime='application/pdf'
    elif format=='xlsx':content=excel_report(result);mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif format=='json':content=json.dumps(result,ensure_ascii=False,indent=2).encode();mime='application/json'
    else:raise ValueError('Поддерживаются PDF, XLSX и JSON')
    response=HttpResponse(content,content_type=mime);response['Content-Disposition']=f'attachment; filename="roboscope-report.{format}"';return response

@api()
def input_template(request):
    wb=Workbook();s=wb.active;s.title='Параметры'
    s.append(['key','Параметр','Значение','Единица','Источник / допущение'])
    d=defaults()
    for f in FIELDS:s.append([f['key'],f['label'],d[f['key']],f['unit'],f['source']])
    for c in ['A','B','C','D','E']:s.column_dimensions[c].width=30 if c!='E' else 75
    out=io.BytesIO();wb.save(out)
    r=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    r['Content-Disposition']='attachment; filename="warehouse-template.xlsx"';return r

@api(('POST',))
def import_inputs(request):
    f=request.FILES.get('file')
    if not f:raise ValueError('Выберите XLSX или CSV')
    if f.size>5*1024*1024:raise ValueError('Размер файла не должен превышать 5 МБ')
    raw=f.read();name=f.name.lower()
    if name.endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(x.file_size for x in z.infolist())>30*1024*1024:raise ValueError('Слишком большой распакованный файл')
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            if 'Склад' in wb.sheetnames:
                s=wb['Склад'];mapping={'area':'C4','active_area':'C5','aisle_width':'C9','payload_kg':'C41','days_year':'C14','peak_factor':'C16','staff':'C28','salary':'C31','payroll_factor':'C32','horizon':'C51'}
                values={k:s[c].value for k,c in mapping.items()};values['hours_day']=s['C13'].value*s['C15'].value
                values['daily_moves']=s['C18'].value+s['C19'].value;values['budget']=s['C50'].value*1000000
                return JsonResponse({'inputs':validate(values,defaults()),'imported':len(values),'notes':['Загружен лист «Склад» организатора. Остальные параметры заполнены нормативами.']})
            sheet=wb.active
            if sheet.max_row>1000:raise ValueError('В шаблоне допускается не более 1 000 строк')
            rows=list(sheet.iter_rows(values_only=True))
        except (zipfile.BadZipFile,KeyError):raise ValueError('Не удалось открыть XLSX')
    elif name.endswith('.csv'):
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:raise ValueError('Сохраните CSV в кодировке UTF-8')
        try:dialect=csv.Sniffer().sniff(text[:2000],delimiters=';,\t')
        except csv.Error:dialect=csv.excel
        rows=list(csv.reader(io.StringIO(text),dialect=dialect))
    else:raise ValueError('Допустимы файлы .xlsx и .csv')
    if not rows:raise ValueError('Файл пуст')
    header=[str(x).strip().lower() for x in rows[0]]
    if 'key' not in header or not ('значение' in header or 'value' in header):raise ValueError('Используйте шаблон: колонки key и Значение')
    ki=header.index('key');vi=header.index('значение' if 'значение' in header else 'value')
    values={}
    for row in rows[1:]:
        if not row or not any(x is not None and x!='' for x in row):continue
        if len(row)<=max(ki,vi):raise ValueError('В строке не хватает столбцов')
        key=str(row[ki]).strip()
        if key in values:raise ValueError('Параметр повторяется: '+key)
        values[key]=row[vi]
    return JsonResponse({'inputs':validate(values,defaults()),'imported':len(values),'notes':['Параметры загружены. Проверьте допущения перед расчётом.']})

@api(('POST',),auth=True,staff=True)
def catalog_upload(request):
    f=request.FILES.get('file')
    if not f or f.size>5*1024*1024:raise ValueError('Выберите CSV до 5 МБ')
    return JsonResponse(import_catalog(f.read()))

@api()
def health(request):
    with connection.cursor() as c:c.execute('SELECT 1');c.fetchone()
    return JsonResponse({'status':'ok','model_version':MODEL_VERSION})

@api()
def openapi(request):
    return JsonResponse(json.loads((settings.BASE_DIR/'docs/openapi.json').read_text()))

@api(auth=True,staff=True)
def worker_health(request):
    heartbeat=ServiceHeartbeat.objects.filter(key='simulation-worker').first()
    age=(timezone.now()-heartbeat.updated_at).total_seconds() if heartbeat else None
    healthy=age is not None and age<90
    return JsonResponse({'status':'ok' if healthy else 'stale','heartbeat_age_seconds':age,
        'pending':SimulationRun.objects.filter(status='pending').count(),
        'running':SimulationRun.objects.filter(status='running').count()},status=200 if healthy else 503)
