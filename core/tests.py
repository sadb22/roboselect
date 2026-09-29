from django.test import override_settings
import io
import json
from copy import deepcopy
from pathlib import Path
from django.test import TestCase, Client
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook, Workbook
from .catalog import import_catalog
from .engine import defaults, validate, calculate, product_dict, applicability, InputError
from .models import Product, Project, ProjectVersion, SimulationRun, Normative
from .simulation import simulate

@override_settings(PASSWORDLESS_DEMO=False)
class ModelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        import_catalog((settings.BASE_DIR/'data/catalog.csv').read_bytes())
        cls.product=product_dict(Product.objects.filter(name__contains='H1500').first())

    def test_catalog_preserves_applications(self):
        self.assertEqual(Product.objects.count(),187)
        self.assertEqual(sum(len(p.observations) for p in Product.objects.all()),223)
        self.assertEqual(len(self.product['applications']),2)

    def test_conflicting_prices_visible(self):
        p=Product.objects.filter(name__contains='А25-1720').first()
        self.assertTrue(product_dict(p)['price_conflict'])
        self.assertEqual(len(p.observations),2)

    def test_import_atomic(self):
        csv='id;Название;компания;Цена изделия\nnew;Test;Test;100\nbroken;Bad;Test;not-price\n'
        with self.assertRaises(ValueError):import_catalog(csv)
        self.assertFalse(Product.objects.filter(source_id='new').exists())

    def test_incompatible_payload_excludes(self):
        v=defaults();v['payload_kg']=1800
        self.assertEqual(applicability(self.product,v)['state'],'excluded')
        with self.assertRaises(InputError):calculate(v,self.product)

    def test_unknown_specs_not_claimed_verified(self):
        p=product_dict(Product.objects.filter(name__contains='H2000').first())
        m=applicability(p,defaults())
        self.assertEqual(m['state'],'review');self.assertTrue(m['missing'])

    def test_bad_inputs_rejected(self):
        for key,value in [('daily_moves',0),('speed','NaN'),('hours_day',25),('horizon',4),('loading_bays',1.5),('floor_type','unknown')]:
            with self.subTest(key=key):
                with self.assertRaises(InputError):validate({key:value})
        with self.assertRaises(InputError):validate({'active_area':30000,'area':10000})
        with self.assertRaises(InputError):validate({'salary':True})
        with self.assertRaises(InputError):validate({'extra_field':1})

    def test_locale_decimal(self):
        self.assertEqual(validate({'robot_price':'2 700 000,00'})['robot_price'],2700000)

    def test_fleet_independent_simple_example(self):
        v=defaults();v.update(distance=30,speed=1,handling_seconds=60,battery_hours=1,charge_minutes=60,
                              daily_moves=1000,hours_day=10,peak_factor=1,utilization=50,reserve_pct=0)
        r=calculate(v,self.product,False)
        # 120 s/cycle; 50% battery availability -> 15/h; target utilization -> 7.5/h; ceil(100/7.5)=14.
        self.assertEqual(r['fleet']['robots'],14)
        self.assertEqual(r['fleet']['per_robot_capacity'],15)

    def test_baseline_uses_annual_fte_once(self):
        r=calculate(defaults(),self.product)
        base=r['scenarios'][0]
        self.assertEqual(base['opex'],25*120000*12*1.302)
        self.assertEqual(base['tco'],234360000)
        self.assertIsNone(base['roi'])

    def test_tco_and_cashflow_reconcile(self):
        r=calculate(defaults(),self.product)
        baseline=r['scenarios'][0]['tco']
        for s in r['scenarios']:
            self.assertAlmostEqual(baseline-s['tco'],s['net_effect'])
            self.assertAlmostEqual(s['cashflow'][-1]['cumulative'],s['net_effect'])
        buy=r['scenarios'][1]
        self.assertEqual(buy['cashflow'][3]['replacement'],17*250000)
        self.assertEqual(buy['cashflow'][4]['replacement'],0)

    def test_released_time_is_not_automatically_savings(self):
        v=defaults();v['cash_realization_pct']=0
        r=calculate(v,self.product)
        self.assertEqual(r['fleet']['saved_fte'],0)
        self.assertLess(r['scenarios'][1]['annual_effect'],0)
        self.assertIsNone(r['scenarios'][1]['payback'])

    def test_zero_capex_service(self):
        v=defaults();v['raas_setup']=0;r=calculate(v,self.product)
        self.assertIsNone(r['scenarios'][2]['roi']);self.assertIsNone(r['scenarios'][2]['payback'])

    def test_manual_robot_count_is_preserved(self):
        v=defaults();v['robots_override']=2;r=calculate(v,self.product)
        self.assertEqual(r['fleet']['robots'],2)
        self.assertTrue(any('ниже расчётного' in x for x in r['warnings']))
        self.assertEqual({row['robots'] for row in r['sensitivity']},{2})

    def test_sensitivity_recomputes_discrete_fleet(self):
        r=calculate(defaults(),self.product)
        rows=[x for x in r['sensitivity'] if x['parameter']=='daily_moves']
        self.assertLess(rows[0]['robots'],rows[2]['robots'])
        self.assertEqual(len(r['sensitivity']),9)

    def test_simulation_deterministic_and_conserves_jobs(self):
        r=calculate(defaults(),self.product,False)
        a=simulate(r,hours=1);b=simulate(r,hours=1)
        self.assertEqual(a,b)
        self.assertEqual(a['arrived'],a['completed']+a['queued']+a['in_progress'])
        self.assertAlmostEqual(sum(a['state_shares'].values()),100,delta=.1)
        for e in a['events']:self.assertGreaterEqual(e['end'],e['start']);self.assertLessEqual(e['end'],a['duration'])

    def test_simulation_detects_bottleneck(self):
        v=defaults();r=calculate(v,self.product,False);bad=simulate(r)
        self.assertFalse(bad['passed']);self.assertGreater(bad['queued'],100)
        v.update(loading_bays=6,corridor_capacity=14)
        better=simulate(calculate(v,self.product,False))
        self.assertGreater(better['throughput'],bad['throughput']*1.5)
        self.assertLess(better['queued'],bad['queued'])

    def test_simulation_limits(self):
        v=defaults();v['robots_override']=81
        with self.assertRaises(ValueError):simulate(calculate(v,self.product,False))

@override_settings(PASSWORDLESS_DEMO=False)
class APITests(TestCase):
    @classmethod
    def setUpTestData(cls):
        import_catalog((settings.BASE_DIR/'data/catalog.csv').read_bytes())
        cls.product=Product.objects.filter(name__contains='H1500').first()
        cls.owner=get_user_model().objects.create_user('owner','', 'Very-Strong-Test-Password42')
        cls.other=get_user_model().objects.create_user('other','', 'Another-Test-Password42')

    def payload(self):return {'inputs':defaults(),'product_id':str(self.product.id),'object_type':'warehouse','name':'Test warehouse'}
    def post(self,url,data):return self.client.post(url,data=json.dumps(data),content_type='application/json')
    def save(self):
        self.client.force_login(self.owner)
        response=self.post('/api/projects/',self.payload());self.assertEqual(response.status_code,201,response.content)
        return response.json()

    def test_guest_calculation(self):
        self.assertEqual(self.client.get('/api/bootstrap/').status_code,200)
        self.assertEqual(self.post('/api/calculate/',self.payload()).status_code,200)
        self.assertEqual(self.post('/api/projects/',self.payload()).status_code,401)

    def test_malformed_product(self):
        for product in [None,'invalid']:
            p=self.payload();p['product_id']=product
            self.assertEqual(self.post('/api/calculate/',p).status_code,400)

    def test_nonwarehouse_not_fake_calculated(self):
        p=self.payload();p['object_type']='hospital'
        self.assertEqual(self.post('/api/calculate/',p).status_code,400)

    def test_project_isolation_all_actions(self):
        p=self.save();id=p['id'];self.client.force_login(self.other)
        self.assertEqual(self.client.get('/api/projects/').json()['projects'],[])
        for method,url in [('get',f'/api/projects/{id}/'),('delete',f'/api/projects/{id}/'),('post',f'/api/projects/{id}/copy/')]:
            self.assertEqual(getattr(self.client,method)(url).status_code,404)
        self.assertEqual(self.post('/api/export/pdf/',{'project_id':id,'revision':1}).status_code,404)

    def test_version_reproduces_original_after_catalog_changes(self):
        p=self.save();id=p['id'];before=self.client.get(f'/api/projects/{id}/').json()['result']
        Product.objects.filter(pk=self.product.pk).update(price=99999999)
        after=self.client.get(f'/api/projects/{id}/?revision=1').json()['result']
        self.assertEqual(before,after)

    def test_concurrent_save_conflict(self):
        p=self.save();payload=self.payload();payload['revision']=1
        u=f"/api/projects/{p['id']}/"
        first=self.client.put(u,json.dumps(payload),content_type='application/json')
        second=self.client.put(u,json.dumps(payload),content_type='application/json')
        self.assertEqual(first.status_code,200);self.assertEqual(second.status_code,409)
        self.assertEqual(ProjectVersion.objects.count(),2)

    def test_copy_and_delete_versions(self):
        p=self.save();copy=self.post(f"/api/projects/{p['id']}/copy/",{});self.assertEqual(copy.status_code,201)
        self.assertEqual(Project.objects.count(),2)
        self.client.delete(f"/api/projects/{p['id']}/")
        self.assertEqual(ProjectVersion.objects.count(),1)

    def test_registration_and_password_hash(self):
        response=self.post('/api/auth/register/',{'username':'new-user','password':'Rocket-Test-Secure-6742'})
        self.assertEqual(response.status_code,404)
        self.assertFalse(get_user_model().objects.filter(username='new-user').exists())

    def test_login_lockout(self):
        for _ in range(8):self.post('/api/auth/login/',{'username':'owner','password':'bad'})
        self.assertEqual(self.post('/api/auth/login/',{'username':'owner','password':'bad'}).status_code,429)

    def test_csrf_required(self):
        c=Client(enforce_csrf_checks=True)
        self.assertEqual(c.post('/api/calculate/',json.dumps(self.payload()),content_type='application/json').status_code,403)

    def test_csv_import_and_bad_input(self):
        good=SimpleUploadedFile('input.csv','key;Значение\ndaily_moves;1500\n'.encode())
        r=self.client.post('/api/import/',{'file':good});self.assertEqual(r.status_code,200,r.content)
        self.assertEqual(r.json()['inputs']['daily_moves'],1500)
        bad=SimpleUploadedFile('input.csv','key;Значение\nspeed;0\n'.encode())
        self.assertEqual(self.client.post('/api/import/',{'file':bad}).status_code,400)

    def test_excel_template_roundtrip(self):
        template=self.client.get('/api/template/');self.assertEqual(template.status_code,200)
        r=self.client.post('/api/import/',{'file':SimpleUploadedFile('template.xlsx',template.content)})
        self.assertEqual(r.status_code,200,r.content);self.assertEqual(r.json()['inputs'],defaults())

    def test_report_formats_and_values(self):
        for fmt in ['pdf','xlsx','json']:
            response=self.post('/api/export/'+fmt+'/',self.payload())
            self.assertEqual(response.status_code,200,response.content[:500])
            if fmt=='pdf':self.assertTrue(response.content.startswith(b'%PDF'))
            if fmt=='xlsx':
                w=load_workbook(io.BytesIO(response.content),data_only=True)
                self.assertEqual(w['Сравнение']['G5'].value,234360000)
                self.assertEqual(w['Сравнение']['A7'].value,'Услуга')

    def test_simulation_run_private(self):
        self.client.force_login(self.owner);run=self.post('/api/simulations/',self.payload())
        self.assertEqual(run.status_code,202,run.content)
        id=run.json()['id'];self.client.force_login(self.other)
        self.assertEqual(self.client.get(f'/api/simulations/{id}/').status_code,404)

    def test_only_one_active_simulation(self):
        self.assertEqual(self.post('/api/simulations/',self.payload()).status_code,202)
        self.assertEqual(self.post('/api/simulations/',self.payload()).status_code,409)

    def test_catalog_import_requires_admin(self):
        self.client.force_login(self.owner)
        self.assertEqual(self.post('/api/catalog/import/',{}).status_code,403)

    def test_api_schema_and_health(self):
        self.assertEqual(self.client.get('/health/').status_code,200)
        schema=self.client.get('/api/schema/')
        self.assertEqual(schema.status_code,200)
        self.assertEqual(schema.json()['openapi'],'3.0.3')

    def test_simulation_summary_is_saved_and_exported(self):
        self.client.force_login(self.owner)
        data=self.payload()
        calc=self.post('/api/calculate/',data).json()
        run=SimulationRun.objects.create(owner=self.owner,status='done',payload={'calculation':calc},result=simulate(calc,42,1))
        data['simulation_id']=str(run.id)
        saved=self.post('/api/projects/',data)
        self.assertEqual(saved.status_code,201,saved.content)
        project=Project.objects.get(id=saved.json()['id'])
        self.assertIn('simulation_summary',project.result)
        self.assertNotIn('events',project.result['simulation_summary'])
        report=self.post('/api/export/xlsx/',{'project_id':str(project.id)})
        wb=load_workbook(io.BytesIO(report.content))
        self.assertIn('Симуляция',wb.sheetnames)
        self.client.force_login(self.other)
        self.assertEqual(self.post('/api/projects/',data).status_code,404)

    def test_changed_parameters_cannot_reuse_simulation(self):
        self.client.force_login(self.owner)
        data=self.payload();calc=self.post('/api/calculate/',data).json()
        run=SimulationRun.objects.create(owner=self.owner,status='done',payload={'calculation':calc},result={'passed':True})
        data['simulation_id']=str(run.id);data['inputs']['daily_moves']=1900
        self.assertEqual(self.post('/api/projects/',data).status_code,400)

    def test_resave_snapshot_keeps_simulation_and_separates_revisions(self):
        self.client.force_login(self.owner)
        data=self.payload();calc=self.post('/api/calculate/',data).json()
        run=SimulationRun.objects.create(owner=self.owner,status='done',payload={'calculation':calc},result=simulate(calc,42,1))
        data['simulation_id']=str(run.id)
        created=self.post('/api/projects/',data).json();id=created['id']
        changed=self.payload();changed['inputs']['daily_moves']=1900;changed['revision']=1
        self.assertEqual(self.client.put(f'/api/projects/{id}/',json.dumps(changed),content_type='application/json').status_code,200)
        saved=self.client.put(f'/api/projects/{id}/',json.dumps({'project_id':id,'calculation_revision':1,'revision':2}),content_type='application/json')
        self.assertEqual(saved.status_code,200,saved.content)
        project=Project.objects.get(id=id)
        self.assertEqual(project.revision,3)
        self.assertEqual(project.result['inputs'],calc['inputs'])
        self.assertIn('simulation_summary',project.result)

    def test_anonymous_simulation_claimed_only_by_same_session(self):
        data=self.payload();created=self.post('/api/simulations/',data).json()
        get_user_model().objects.create_user(username='claim-user',password='Claim-Session-Test-9841')
        response=self.post('/api/auth/login/',{'username':'claim-user','password':'Claim-Session-Test-9841'})
        self.assertEqual(response.status_code,200,response.content)
        run=SimulationRun.objects.get(id=created['id'])
        self.assertEqual(run.owner.username,'claim-user')
